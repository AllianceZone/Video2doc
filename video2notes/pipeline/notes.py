"""
Notes-specific pipeline stages: grouping frames into topically-coherent
sections, grouping sections further for real (non-degenerate) summaries via
the configured LLM provider, and building the final Word/PowerPoint artifacts.
"""

import logging
import re
from pathlib import Path
from typing import Dict, List, Optional

from llm.base import LLMProvider

logger = logging.getLogger("video2notes.pipeline.notes")


# ---------------------------------------------------------------------------
# Semantic chunking (frames -> sections)
# ---------------------------------------------------------------------------

def _frame_start(frame_path: Path) -> float:
    stamp = frame_path.stem.replace("frame_", "")
    h, m, s = (int(x) for x in stamp.split("-"))
    return h * 3600 + m * 60 + s


def _segments_between(segments: List[Dict], start: float, end: float) -> List[Dict]:
    return [seg for seg in segments if start <= seg["start"] < end]


def compute_sections(frame_paths: List[Path], segments: List[Dict],
                      similarity_threshold: float = 0.15, max_chunk_seconds: float = 60.0,
                      max_frames_per_chunk: int = 4) -> List[Dict]:
    """Groups consecutive frames whose transcript is topically similar (TF-IDF
    cosine) and close in time into one 'section'. Returns a list of
    {"start", "end", "frames": [Path,...], "segments": [Dict,...]}."""
    frame_paths = sorted(frame_paths, key=_frame_start)
    if not frame_paths:
        return []

    frame_texts = []
    for i, f in enumerate(frame_paths):
        start = _frame_start(f)
        end = _frame_start(frame_paths[i + 1]) if i + 1 < len(frame_paths) else start + 3600
        frame_texts.append(" ".join(s["text"] for s in _segments_between(segments, start, end)))

    similarity = _pairwise_similarity(frame_texts)

    chunks: List[List[Path]] = [[frame_paths[0]]]
    for i in range(1, len(frame_paths)):
        cur = chunks[-1]
        chunk_start = _frame_start(cur[0])
        candidate_start = _frame_start(frame_paths[i])
        time_ok = (candidate_start - chunk_start) <= max_chunk_seconds
        size_ok = len(cur) < max_frames_per_chunk
        sim_ok = similarity[i - 1][i] >= similarity_threshold if similarity is not None else False
        if time_ok and size_ok and sim_ok:
            cur.append(frame_paths[i])
        else:
            chunks.append([frame_paths[i]])

    result = []
    for i, chunk_frames in enumerate(chunks):
        start = _frame_start(chunk_frames[0])
        end = _frame_start(chunks[i + 1][0]) if i + 1 < len(chunks) else start + 3600
        chunk_segments = _segments_between(segments, start, end)
        result.append({"start": start, "end": end, "frames": chunk_frames, "segments": chunk_segments})

    logger.info(f"Grouped {len(frame_paths)} frames into {len(result)} sections")
    return result


def representative_images(section: Dict) -> List[Path]:
    frames = section["frames"]
    if len(frames) == 1 or frames[0] == frames[-1]:
        return [frames[0]]
    return [frames[0], frames[-1]]


def _pairwise_similarity(texts: List[str]):
    if len(texts) < 2:
        return None
    try:
        from sklearn.feature_extraction.text import TfidfVectorizer
        from sklearn.metrics.pairwise import cosine_similarity
    except ImportError:
        return None
    if all(not t.strip() for t in texts):
        return None
    try:
        vectorizer = TfidfVectorizer(stop_words="english")
        tfidf = vectorizer.fit_transform([t if t.strip() else " " for t in texts])
        return cosine_similarity(tfidf)
    except ValueError:
        return None


# ---------------------------------------------------------------------------
# Summary grouping (sections -> summary groups) + LLM call
# ---------------------------------------------------------------------------

def group_sections_for_summary(sections: List[Dict], group_size: int = 6,
                                max_group_seconds: float = 180.0) -> List[Dict]:
    """Combines several sections together before summarizing -- summarizing
    one tiny section on its own has nothing to compress. Returns
    {"start", "end", "sections": [...], "segments": [...]}."""
    if not sections:
        return []
    groups: List[List[Dict]] = [[sections[0]]]
    for s in sections[1:]:
        cur = groups[-1]
        span = s["end"] - cur[0]["start"]
        if len(cur) >= group_size or span > max_group_seconds:
            groups.append([s])
        else:
            cur.append(s)

    result = []
    for g in groups:
        segments = [seg for sec in g for seg in sec["segments"]]
        result.append({"start": g[0]["start"], "end": g[-1]["end"], "sections": g, "segments": segments})
    return result


def summarize_group(group: Dict, provider: LLMProvider, style: str = "action_items",
                     max_summary_sentences: int = 3, max_key_points: int = 6) -> Dict:
    text = " ".join(seg["text"].strip() for seg in group["segments"] if seg["text"].strip())
    result = provider.summarize(text, style=style, max_summary_sentences=max_summary_sentences,
                                 max_key_points=max_key_points)
    return {"summary": result["summary"], "key_points": result["key_points"]}


# ---------------------------------------------------------------------------
# Document generation
# ---------------------------------------------------------------------------

def _hhmmss(seconds: float) -> str:
    total = int(round(seconds))
    h, rem = divmod(total, 3600)
    m, s = divmod(rem, 60)
    return f"{h:02d}:{m:02d}:{s:02d}"


def build_audio_docx(segments: List[Dict], title: str, output_path: Path,
                      summary: str = "", key_points: Optional[List[str]] = None) -> Path:
    """For audio-only input (no frames): title, summary, key points, full
    timestamped transcript."""
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    t = doc.add_heading(title, level=0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("Auto-generated audio transcript report").alignment = WD_ALIGN_PARAGRAPH.CENTER

    if summary:
        doc.add_heading("Summary", level=1)
        doc.add_paragraph(summary)
    if key_points:
        doc.add_heading("Key Points", level=1)
        for kp in key_points:
            doc.add_paragraph(kp, style="List Bullet")

    doc.add_page_break()
    doc.add_heading("Full Transcript", level=1)
    for seg in segments:
        text = seg["text"].strip()
        if not text:
            continue
        p = doc.add_paragraph()
        run_ts = p.add_run(f"[{_hhmmss(seg['start'])}]  ")
        run_ts.bold = True
        p.add_run(text)
        p.paragraph_format.space_after = Pt(8)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    return output_path


def build_audio_pptx(segments: List[Dict], title: str, output_path: Path,
                      summary: str = "", key_points: Optional[List[str]] = None) -> Path:
    """For audio-only input: title, summary, key points, paginated transcript slides."""
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN
    from pptx.dml.color import RGBColor

    COLOR_TITLE = RGBColor(0x1B, 0x2A, 0x41)
    COLOR_BODY = RGBColor(0x33, 0x38, 0x40)
    COLOR_CARD_BG = RGBColor(0xF4, 0xF6, 0xF5)

    def truncate_text(text: str, max_chars: int) -> str:
        if len(text) <= max_chars:
            return text
        return text[:max_chars].rsplit(" ", 1)[0].rstrip(",.;: ") + "…"

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    def blank_slide():
        slide = prs.slides.add_slide(blank_layout)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        return slide

    def add_textbox(slide, left, top, width, height, text, size=18, bold=False, color=COLOR_BODY):
        box = slide.shapes.add_textbox(left, top, width, height)
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = text
        run.font.size, run.font.bold, run.font.color.rgb = Pt(size), bold, color
        return box

    def add_bullets(slide, left, top, width, height, items, size=14):
        box = slide.shapes.add_textbox(left, top, width, height)
        tf = box.text_frame
        tf.word_wrap = True
        for i, item in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(8)
            run = p.add_run()
            run.text = f"•  {item}"
            run.font.size, run.font.color.rgb = Pt(size), COLOR_BODY

    def add_card(slide, left, top, width, height):
        from pptx.enum.shapes import MSO_SHAPE
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = COLOR_CARD_BG
        shape.line.fill.background()
        shape.shadow.inherit = False

    slide = blank_slide()
    add_textbox(slide, Inches(1), Inches(2.7), Inches(11.3), Inches(1.4), title, size=40, bold=True, color=COLOR_TITLE)
    add_textbox(slide, Inches(1), Inches(3.9), Inches(11.3), Inches(0.6), "Audio transcript report", size=18)

    if summary:
        slide = blank_slide()
        add_textbox(slide, Inches(0.7), Inches(0.6), Inches(11.9), Inches(0.8), "Summary", size=32, bold=True, color=COLOR_TITLE)
        add_card(slide, Inches(0.7), Inches(1.6), Inches(11.9), Inches(5.2))
        add_textbox(slide, Inches(1.1), Inches(2.0), Inches(11.1), Inches(4.4), truncate_text(summary, 1400), size=18)
    if key_points:
        slide = blank_slide()
        add_textbox(slide, Inches(0.7), Inches(0.6), Inches(11.9), Inches(0.8), "Key Points", size=32, bold=True, color=COLOR_TITLE)
        add_card(slide, Inches(0.7), Inches(1.6), Inches(11.9), Inches(5.2))
        add_bullets(slide, Inches(1.1), Inches(2.0), Inches(11.1), Inches(4.4), key_points[:10], size=15)

    items = [f"[{_hhmmss(seg['start'])}]  {seg['text'].strip()}" for seg in segments if seg["text"].strip()]
    pages, current, chars = [], [], 0
    for item in items:
        if current and (len(current) >= 10 or chars + len(item) > 1500):
            pages.append(current)
            current, chars = [], 0
        current.append(item)
        chars += len(item)
    if current:
        pages.append(current)

    for i, page_items in enumerate(pages, 1):
        slide = blank_slide()
        add_textbox(slide, Inches(0.7), Inches(0.5), Inches(11.9), Inches(0.7),
                    f"Transcript ({i}/{len(pages)})", size=26, bold=True, color=COLOR_TITLE)
        add_card(slide, Inches(0.7), Inches(1.35), Inches(11.9), Inches(5.7))
        font_size = 15 if sum(len(t) for t in page_items) <= 500 else 13
        add_bullets(slide, Inches(1.1), Inches(1.7), Inches(11.1), Inches(5.1), page_items, size=font_size)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))
    return output_path


def build_docx(groups: List[Dict], video_title: str, output_path: Path,
                overall_summary: str = "", overall_key_points: Optional[List[str]] = None,
                image_width_inches: float = 5.5) -> Path:
    """Renders the 4-part report: for each summary group, its takeaways, then
    for each nested section -- image(s), an audio-note hyperlink, transcript.
    Each section dict must additionally carry "images" (List[Path]),
    "audio_clip" (Path or None), "transcript_text" (str)."""
    from docx import Document
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.oxml.ns import qn
    from docx.oxml.shared import OxmlElement
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    def add_hyperlink(paragraph, text: str, target_path: str):
        part = paragraph.part
        r_id = part.relate_to(target_path, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), r_id)
        run = OxmlElement("w:r")
        rPr = OxmlElement("w:rPr")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "2E6B5E")
        rPr.append(color)
        underline = OxmlElement("w:u")
        underline.set(qn("w:val"), "single")
        rPr.append(underline)
        run.append(rPr)
        t = OxmlElement("w:t")
        t.text = text
        run.append(t)
        hyperlink.append(run)
        paragraph._p.append(hyperlink)

    total_sections = sum(len(g["sections"]) for g in groups)
    doc = Document()
    title = doc.add_heading(video_title, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(f"Auto-generated report  •  {total_sections} sections").alignment = WD_ALIGN_PARAGRAPH.CENTER

    if overall_summary:
        doc.add_heading("Overall Summary", level=1)
        doc.add_paragraph(overall_summary)
    if overall_key_points:
        doc.add_heading("Overall Key Points", level=1)
        for kp in overall_key_points:
            doc.add_paragraph(kp, style="List Bullet")

    doc.add_page_break()
    counter = 0
    for gi, group in enumerate(groups, 1):
        if group.get("summary") or group.get("key_points"):
            doc.add_heading(f"[{_hhmmss(group['start'])} – {_hhmmss(group['end'])}]", level=1)
            if group.get("summary"):
                doc.add_paragraph(group["summary"]).paragraph_format.space_after = Pt(4)
            if group.get("key_points"):
                p = doc.add_paragraph()
                run = p.add_run("Key Takeaways")
                run.bold = True
                run.font.color.rgb = RGBColor(0x2E, 0x6B, 0x5E)
                for kp in group["key_points"]:
                    doc.add_paragraph(kp, style="List Bullet")
            doc.add_paragraph("").paragraph_format.space_after = Pt(2)

        for section in group["sections"]:
            counter += 1
            doc.add_heading(f"Section {counter}  ·  [{_hhmmss(section['start'])} – {_hhmmss(section['end'])}]",
                             level=2)
            for img in section.get("images", []):
                try:
                    doc.add_picture(str(img), width=Inches(image_width_inches))
                except Exception as e:
                    logger.warning(f"Could not insert image {img}: {e}")
            audio_clip = section.get("audio_clip")
            if audio_clip:
                p = doc.add_paragraph()
                add_hyperlink(p, f"🔊 Play audio note ({Path(audio_clip).name})",
                               f"audio_notes/{Path(audio_clip).name}")
                p.paragraph_format.space_after = Pt(6)
            text = section.get("transcript_text", "")
            if text:
                doc.add_paragraph(text).paragraph_format.space_after = Pt(8)
            else:
                doc.add_paragraph("(no speech detected in this section)").paragraph_format.space_after = Pt(8)

        if gi < len(groups):
            doc.add_paragraph("").paragraph_format.space_after = Pt(4)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    return output_path


def build_pptx(groups: List[Dict], video_title: str, output_path: Path,
               overall_summary: str = "", overall_key_points: Optional[List[str]] = None) -> Path:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.dml.color import RGBColor
    from PIL import Image

    COLOR_TITLE = RGBColor(0x1B, 0x2A, 0x41)
    COLOR_BODY = RGBColor(0x33, 0x38, 0x40)
    COLOR_ACCENT = RGBColor(0x2E, 0x6B, 0x5E)
    COLOR_CARD_BG = RGBColor(0xF4, 0xF6, 0xF5)
    COLOR_MUTED = RGBColor(0x76, 0x7E, 0x87)
    SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)

    def split_sentences(text: str) -> List[str]:
        text = text.strip()
        if not text:
            return []
        return [s.strip() for s in re.split(r'(?<=[.!?])\s+(?=[A-Z0-9"\'])', text) if s.strip()]

    def truncate_text(text: str, max_chars: int) -> str:
        if len(text) <= max_chars:
            return text
        return text[:max_chars].rsplit(" ", 1)[0].rstrip(",.;: ") + "…"

    def fit_bullets(texts: List[str], max_chars: int, max_items: int) -> List[str]:
        fitted, total = [], 0
        for t in texts:
            if len(fitted) >= max_items or total + len(t) > max_chars:
                remaining = len(texts) - len(fitted)
                if remaining > 0:
                    fitted.append(f"(+{remaining} more — see the Word doc for the full transcript)")
                break
            fitted.append(t)
            total += len(t)
        return fitted

    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    blank_layout = prs.slide_layouts[6]

    def blank_slide():
        slide = prs.slides.add_slide(blank_layout)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        return slide

    def add_textbox(slide, left, top, width, height, text, size=18, bold=False,
                     color=COLOR_BODY, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
        box = slide.shapes.add_textbox(left, top, width, height)
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        for attr in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, attr, 0)
        p = tf.paragraphs[0]
        p.alignment = align
        run = p.add_run()
        run.text = text
        run.font.size, run.font.bold, run.font.color.rgb = Pt(size), bold, color
        return box

    def add_bullets(slide, left, top, width, height, items, size=16, color=COLOR_BODY, space_after=10):
        box = slide.shapes.add_textbox(left, top, width, height)
        tf = box.text_frame
        tf.word_wrap = True
        for attr in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, attr, 0)
        for i, item in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(space_after)
            run = p.add_run()
            run.text = f"•  {item}"
            run.font.size, run.font.color.rgb = Pt(size), color

    def add_card(slide, left, top, width, height, fill=COLOR_CARD_BG):
        from pptx.enum.shapes import MSO_SHAPE
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
        shape.line.fill.background()
        shape.shadow.inherit = False
        return shape

    # Title slide
    slide = blank_slide()
    total_sections = sum(len(g["sections"]) for g in groups)
    add_textbox(slide, Inches(1), Inches(2.7), Inches(11.3), Inches(1.4), video_title, size=40, bold=True,
                color=COLOR_TITLE)
    add_textbox(slide, Inches(1), Inches(3.9), Inches(11.3), Inches(0.6),
                f"{total_sections} sections  •  auto-generated report", size=18, color=COLOR_MUTED)

    # Overall summary / key points slides
    if overall_summary:
        slide = blank_slide()
        add_textbox(slide, Inches(0.7), Inches(0.6), Inches(11.9), Inches(0.8), "Summary", size=32, bold=True,
                    color=COLOR_TITLE)
        add_card(slide, Inches(0.7), Inches(1.6), Inches(11.9), Inches(5.2))
        add_textbox(slide, Inches(1.1), Inches(2.0), Inches(11.1), Inches(4.4),
                    truncate_text(overall_summary, 1400), size=18, color=COLOR_BODY)
    if overall_key_points:
        slide = blank_slide()
        add_textbox(slide, Inches(0.7), Inches(0.6), Inches(11.9), Inches(0.8), "Key Points", size=32, bold=True,
                    color=COLOR_TITLE)
        add_card(slide, Inches(0.7), Inches(1.6), Inches(11.9), Inches(5.2))
        add_bullets(slide, Inches(1.1), Inches(2.0), Inches(11.1), Inches(4.4),
                    fit_bullets(overall_key_points, 950, 8), size=15)

    # Per-group takeaways + per-section slides
    for group in groups:
        summary, key_points = group.get("summary", ""), group.get("key_points", [])
        if summary or key_points:
            slide = blank_slide()
            add_textbox(slide, Inches(0.7), Inches(0.5), Inches(11.9), Inches(0.7),
                        f"[{_hhmmss(group['start'])} – {_hhmmss(group['end'])}]", size=26, bold=True,
                        color=COLOR_TITLE)
            if summary and key_points:
                summary_h = Inches(2.15)
                add_card(slide, Inches(0.7), Inches(1.35), Inches(11.9), summary_h)
                add_textbox(slide, Inches(1.1), Inches(1.6), Inches(11.1), summary_h - Inches(0.5),
                            truncate_text(summary, 650), size=14, color=COLOR_BODY)
                kp_top = Inches(1.35) + summary_h + Inches(0.2)
                kp_h = Inches(7.5) - kp_top - Inches(0.4)
                add_card(slide, Inches(0.7), kp_top, Inches(11.9), kp_h)
                add_bullets(slide, Inches(1.1), kp_top + Inches(0.25), Inches(11.1), kp_h - Inches(0.4),
                            fit_bullets(key_points, 650, 6), size=13)
            elif summary:
                add_card(slide, Inches(0.7), Inches(1.35), Inches(11.9), Inches(5.7))
                add_textbox(slide, Inches(1.1), Inches(1.7), Inches(11.1), Inches(5.1),
                            truncate_text(summary, 1200), size=16, color=COLOR_BODY)
            else:
                add_card(slide, Inches(0.7), Inches(1.35), Inches(11.9), Inches(5.7))
                add_bullets(slide, Inches(1.1), Inches(1.7), Inches(11.1), Inches(5.1),
                            fit_bullets(key_points, 1300, 10), size=15)

        for section in group["sections"]:
            slide = blank_slide()
            add_textbox(slide, Inches(0.6), Inches(0.35), Inches(11.5), Inches(0.55),
                        f"[{_hhmmss(section['start'])} – {_hhmmss(section['end'])}]", size=19, bold=True,
                        color=COLOR_ACCENT)

            images = section.get("images", [])
            img_left, img_top = Inches(0.6), Inches(1.0)
            img_area_w, img_area_h = Inches(6.4), Inches(6.05)
            if images:
                n = len(images)
                slot_h = img_area_h / n if n > 1 else img_area_h
                for idx, img in enumerate(images):
                    try:
                        with Image.open(img) as im:
                            iw, ih = im.size
                        gap = Inches(0.15) if n > 1 else 0
                        box_w, box_h = img_area_w, int(slot_h) - gap
                        if (iw / ih) > (box_w / box_h):
                            draw_w = box_w
                            draw_h = Inches(box_w / (iw / ih) / Inches(1))
                        else:
                            draw_h = box_h
                            draw_w = Inches(box_h * (iw / ih) / Inches(1))
                        left = img_left + (box_w - draw_w) / 2
                        top = img_top + int(idx * slot_h) + (box_h - draw_h) / 2
                        slide.shapes.add_picture(str(img), left, top, width=draw_w, height=draw_h)
                    except Exception as e:
                        logger.warning(f"Could not insert image {img}: {e}")

            text_left = Inches(7.3)
            audio_path = section.get("audio_clip")
            if audio_path:
                try:
                    slide.shapes.add_movie(str(audio_path), text_left, Inches(1.0), Inches(0.55), Inches(0.55),
                                            poster_frame_image=None, mime_type="audio/mpeg")
                    add_textbox(slide, text_left + Inches(0.75), Inches(1.08), Inches(3.5), Inches(0.4),
                                "Audio note", size=12, color=COLOR_MUTED)
                except Exception as e:
                    logger.warning(f"Could not embed audio {audio_path}: {e}")

            card_top = Inches(1.75)
            add_card(slide, text_left, card_top, Inches(5.45), Inches(5.3))
            content = split_sentences(section.get("transcript_text", ""))
            if content:
                fitted = fit_bullets(content, 1150, 11)
                fs = 14 if sum(len(t) for t in fitted) <= 700 else (13 if sum(len(t) for t in fitted) <= 950 else 11)
                add_bullets(slide, text_left + Inches(0.3), card_top + Inches(0.3), Inches(4.85), Inches(4.7),
                            fitted, size=fs, space_after=8)
            else:
                add_textbox(slide, text_left + Inches(0.3), card_top + Inches(0.3), Inches(4.85), Inches(1.0),
                            "(no speech detected in this section)", size=13, color=COLOR_MUTED)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))
    return output_path
